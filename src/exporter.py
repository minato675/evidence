from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def export_results(
    results: list[dict[str, Any]],
    metadata: pd.DataFrame,
    skipped: list[dict[str, Any]],
    output_dir: str | Path,
    schema: dict[str, Any] | None = None,
) -> None:
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)

    _safe_to_csv(metadata, output_path / "metadata.csv")
    _safe_to_csv(pd.DataFrame(skipped), output_path / "skipped_papers.csv")

    jsonl_path = _fallback_path_if_locked(output_path / "extracted_database.jsonl")
    with jsonl_path.open("w", encoding="utf-8") as handle:
        for result in results:
            handle.write(json.dumps(result, ensure_ascii=False) + "\n")

    if any("records" in result for result in results):
        rows, groups = _flatten_data_point_results(results, schema)
        frame = pd.DataFrame(rows)
        _safe_to_csv(frame, output_path / "extracted_database.csv")
        _safe_to_xlsx(frame, groups, output_path / "extracted_database.xlsx")
    else:
        _safe_to_csv(pd.DataFrame([_flatten_result(result) for result in results]), output_path / "extracted_database.csv")


def _flatten_result(result: dict[str, Any]) -> dict[str, Any]:
    row = {
        "paper_id": result.get("paper_id"),
        "title": result.get("title"),
        "doi": result.get("doi"),
    }
    for name, entry in (result.get("extracted_fields") or {}).items():
        value = entry.get("value") if isinstance(entry, dict) else entry
        row[name] = json.dumps(value, ensure_ascii=False) if isinstance(value, list) else value
        if isinstance(entry, dict):
            row[f"{name}_unit"] = entry.get("unit")
            row[f"{name}_evidence_text"] = entry.get("evidence_text")
            row[f"{name}_source_chunk_id"] = entry.get("source_chunk_id")
            row[f"{name}_validation_status"] = entry.get("validation_status")
    return row


def _safe_to_csv(frame: pd.DataFrame, path: Path) -> Path:
    target = _fallback_path_if_locked(path)
    frame.to_csv(target, index=False, encoding="utf-8-sig")
    return target


def _safe_to_xlsx(frame: pd.DataFrame, groups: dict[str, str], path: Path) -> Path:
    target = _fallback_path_if_locked(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "DB"

    columns = list(frame.columns)
    for column_index, column in enumerate(columns, start=1):
        ws.cell(1, column_index, groups.get(column, "")).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, column_index).fill = PatternFill("solid", fgColor=_group_color(groups.get(column, "")))
        ws.cell(1, column_index).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(2, column_index, column).font = Font(bold=True)
        ws.cell(2, column_index).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, row in enumerate(frame.to_dict("records"), start=3):
        for column_index, column in enumerate(columns, start=1):
            ws.cell(row_index, column_index, row.get(column))
            ws.cell(row_index, column_index).alignment = Alignment(vertical="top", wrap_text=True)

    _merge_group_headers(ws, columns, groups)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = ws.dimensions

    for column_index, column in enumerate(columns, start=1):
        letter = get_column_letter(column_index)
        if column in {"Title", "Evidence text", "Notes"}:
            ws.column_dimensions[letter].width = 42
        elif column in {"DOI", "Source chunk"}:
            ws.column_dimensions[letter].width = 24
        elif column in {"Composition", "Material system", "Synthesis method", "Test method"}:
            ws.column_dimensions[letter].width = 22
        else:
            ws.column_dimensions[letter].width = 14

    for row in ws.iter_rows(min_row=3):
        ws.row_dimensions[row[0].row].height = 36

    wb.save(target)
    return target


def _flatten_data_point_results(
    results: list[dict[str, Any]],
    schema: dict[str, Any] | None = None,
) -> tuple[list[dict[str, Any]], dict[str, str]]:
    rows: list[dict[str, Any]] = []
    groups = {
        "Index": "Index",
        "Title": "Meta data",
        "Year": "Meta data",
        "DOI": "Meta data",
    }
    schema_fields = list((schema or {}).get("fields") or [])
    if not schema_fields:
        field_names: list[str] = []
        for result in results:
            for record in result.get("records", []):
                for name in (record.get("fields") or {}):
                    if name not in field_names:
                        field_names.append(name)
        schema_fields = [{"name": name, "group": "Extracted data"} for name in field_names]

    field_columns: list[tuple[str, str]] = []
    for field in schema_fields:
        name = str(field.get("name") or "")
        if not name:
            continue
        label = str(field.get("label") or _field_label(name))
        field_columns.append((name, label))
        groups[label] = str(field.get("group") or "Extracted data")

    evidence_columns = {
        "Evidence text": "Evidence",
        "Source chunk": "Evidence",
        "Source pages": "Evidence",
        "Confidence": "Evidence",
        "Validation status": "Evidence",
    }
    groups.update(evidence_columns)
    index = 1
    for result in results:
        for record in result.get("records", []):
            fields = record.get("fields", {})
            row = {
                "Index": index,
                "Title": result.get("title"),
                "Year": result.get("year"),
                "DOI": result.get("doi"),
            }
            row.update({label: fields.get(name) for name, label in field_columns})
            row.update(
                {
                    "Evidence text": record.get("evidence_text"),
                    "Source chunk": record.get("source_chunk_id"),
                    "Source pages": record.get("source_page_range"),
                    "Confidence": record.get("confidence"),
                    "Validation status": record.get("validation_status"),
                }
            )
            rows.append(row)
            index += 1
    return rows, groups


def _field_label(name: str) -> str:
    return name.replace("_", " ").strip().capitalize()


def _merge_group_headers(ws, columns: list[str], groups: dict[str, str]) -> None:
    start = 1
    while start <= len(columns):
        group = groups.get(columns[start - 1], "")
        end = start
        while end < len(columns) and groups.get(columns[end], "") == group:
            end += 1
        if end > start:
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        start = end + 1


def _group_color(group: str) -> str:
    return {
        "Index": "595959",
        "Meta data": "4472C4",
        "Material": "70AD47",
        "Processing / Sample": "A64D79",
        "Mechanical Property": "ED7D31",
        "Notes": "7F7F7F",
        "Evidence": "2F75B5",
        "Extracted data": "5B9BD5",
    }.get(group, "595959")


def _fallback_path_if_locked(path: Path) -> Path:
    try:
        with path.open("a", encoding="utf-8"):
            return path
    except PermissionError:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return path.with_name(f"{path.stem}_{stamp}{path.suffix}")
